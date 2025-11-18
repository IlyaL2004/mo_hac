import json
import os
from datetime import datetime
from kafka import KafkaProducer
from kafka.errors import KafkaError, KafkaTimeoutError
import hashlib
import time
import psycopg2
from fastapi import FastAPI, HTTPException
from contextlib import asynccontextmanager
import asyncio
from pydantic import BaseModel
from typing import Optional
import webbrowser
import threading
import requests
from fastapi.responses import HTMLResponse
import pandas as pd
import io
import csv
from openpyxl import load_workbook
import docx
import PyPDF2
import chardet
import tempfile
import uuid
from fastapi.middleware.cors import CORSMiddleware



# Модели данных для API
class FileProcessRequest(BaseModel):
    file_path: str
    inn: str


class ApiDownloadRequest(BaseModel):
    api_url: str
    inn: str
    file_name: Optional[str] = None
    file_extension: Optional[str] = "json"


class ProcessResponse(BaseModel):
    status: str
    message: str
    inn: Optional[str] = None
    file_name: Optional[str] = None


class FileContentResponse(BaseModel):
    status: str
    file_name: str
    file_type: str
    content: Optional[dict] = None
    text_content: Optional[str] = None
    error: Optional[str] = None


class ApiDownloadResponse(BaseModel):
    status: str
    message: str
    inn: str
    file_name: str
    api_url: str
    file_size: Optional[int] = None
    download_time: Optional[float] = None


# Класс KafkaConnectFileProducer с расширенной обработкой файлов
class KafkaConnectFileProducer:
    def __init__(self, bootstrap_servers='localhost:9092'):
        self.producer = KafkaProducer(
            bootstrap_servers=bootstrap_servers,
            batch_size=32768,
            linger_ms=100,
            compression_type=None,
            value_serializer=lambda v: json.dumps(v, default=str).encode('utf-8'),
            key_serializer=lambda v: str(v).encode('utf-8') if v else None,
            acks='all',
            request_timeout_ms=60000,
            retries=10,
            retry_backoff_ms=1000,
            max_block_ms=120000,
            buffer_memory=67108864,
        )
        self.global_send_counter = 0

    def get_next_send_id(self):
        """Возвращает следующий глобальный номер отправки"""
        self.global_send_counter += 1
        return self.global_send_counter

    def create_message_with_schema(self, payload):
        """Создает сообщение с JSON-схемой для Kafka Connect"""
        schema = {
            "type": "struct",
            "fields": [
                {"type": "string", "optional": False, "field": "inn"},
                {"type": "string", "optional": False, "field": "file_name"},
                {"type": "string", "optional": False, "field": "file_extension"},
                {"type": "string", "optional": False, "field": "file_data"},
                {"type": "string", "optional": False, "field": "file_hash"},
                {"type": "int32", "optional": False, "field": "file_size"},
                {"type": "string", "optional": False, "field": "processing_date"},
                {"type": "string", "optional": False, "field": "timestamp"},
                {"type": "string", "optional": False, "field": "path_s3"},
                {"type": "string", "optional": True, "field": "parsed_content"},
                {"type": "string", "optional": True, "field": "file_type"}
            ],
            "optional": False,
            "name": "file_chunk"
        }

        return {
            "schema": schema,
            "payload": payload
        }

    def detect_encoding(self, file_bytes):
        """Определяет кодировку файла"""
        try:
            result = chardet.detect(file_bytes)
            encoding = result['encoding'] if result['encoding'] else 'utf-8'
            confidence = result['confidence']
            print(f"Определена кодировка: {encoding} (уверенность: {confidence:.2%})")
            return encoding
        except:
            return 'utf-8'

    def extract_csv_content(self, file_data, file_name):
        """Извлекает содержимое CSV файла"""
        try:
            encoding = self.detect_encoding(file_data)
            text_content = file_data.decode(encoding)

            csv_reader = csv.reader(io.StringIO(text_content))
            rows = list(csv_reader)

            content = {
                "headers": rows[0] if rows else [],
                "rows": rows[1:] if len(rows) > 1 else [],
                "row_count": len(rows) - 1 if len(rows) > 1 else 0,
                "encoding": encoding
            }

            print(f"CSV файл {file_name}: {content['row_count']} строк")
            return content
        except Exception as e:
            print(f"Ошибка при обработке CSV: {e}")
            return None

    def extract_excel_content(self, file_data, file_name):
        """Извлекает содержимое Excel файла"""
        try:
            excel_file = io.BytesIO(file_data)
            excel_data = pd.ExcelFile(excel_file)

            content = {
                "sheets": excel_data.sheet_names,
                "data": {}
            }

            for sheet_name in excel_data.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                content["data"][sheet_name] = {
                    "headers": df.columns.tolist(),
                    "rows": df.fillna('').astype(str).values.tolist(),
                    "row_count": len(df),
                    "columns_count": len(df.columns)
                }

            print(f"Excel файл {file_name}: {len(content['sheets'])} листов")
            return content
        except Exception as e:
            print(f"Ошибка при обработке Excel: {e}")
            return None

    def extract_text_content(self, file_data, file_name):
        """Извлекает содержимое текстового файла"""
        try:
            encoding = self.detect_encoding(file_data)
            text_content = file_data.decode(encoding)

            lines = text_content.split('\n')
            content = {
                "line_count": len(lines),
                "char_count": len(text_content),
                "encoding": encoding,
                "preview": text_content[:1000]  # первые 1000 символов
            }

            print(f"Текстовый файл {file_name}: {content['line_count']} строк")
            return content
        except Exception as e:
            print(f"Ошибка при обработке текстового файла: {e}")
            return None

    def extract_json_content(self, file_data, file_name):
        """Извлекает содержимое JSON файла"""
        try:
            encoding = self.detect_encoding(file_data)
            text_content = file_data.decode(encoding)

            json_data = json.loads(text_content)
            content = {
                "data": json_data,
                "type": type(json_data).__name__,
                "encoding": encoding
            }

            print(f"JSON файл {file_name}: успешно распарсен")
            return content
        except Exception as e:
            print(f"Ошибка при обработке JSON: {e}")
            return None

    def extract_pdf_content(self, file_data, file_name):
        """Извлекает содержимое PDF файла"""
        try:
            pdf_file = io.BytesIO(file_data)
            pdf_reader = PyPDF2.PdfReader(pdf_file)

            text_content = ""
            for page in pdf_reader.pages:
                text_content += page.extract_text() + "\n"

            content = {
                "page_count": len(pdf_reader.pages),
                "text_content": text_content.strip(),
                "char_count": len(text_content)
            }

            print(f"PDF файл {file_name}: {content['page_count']} страниц")
            return content
        except Exception as e:
            print(f"Ошибка при обработке PDF: {e}")
            return None

    def extract_word_content(self, file_data, file_name):
        """Извлекает содержимое Word документа"""
        try:
            doc_file = io.BytesIO(file_data)
            doc = docx.Document(doc_file)

            text_content = ""
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content += paragraph.text + "\n"

            content = {
                "paragraph_count": len(doc.paragraphs),
                "text_content": text_content.strip(),
                "char_count": len(text_content)
            }

            print(f"Word файл {file_name}: {content['paragraph_count']} параграфов")
            return content
        except Exception as e:
            print(f"Ошибка при обработке Word документа: {e}")
            return None

    def extract_file_content(self, file_path):
        """Основная функция для извлечения содержимого файла"""
        file_name = os.path.basename(file_path)
        file_extension = file_name.split('.')[-1].lower() if '.' in file_name else 'bin'

        with open(file_path, 'rb') as f:
            file_data = f.read()

        file_processors = {
            'csv': self.extract_csv_content,
            'xlsx': self.extract_excel_content,
            'xls': self.extract_excel_content,
            'txt': self.extract_text_content,
            'log': self.extract_text_content,
            'json': self.extract_json_content,
            'pdf': self.extract_pdf_content,
            'docx': self.extract_word_content,
            'doc': self.extract_word_content,
        }

        # Текстовые форматы
        text_extensions = ['xml', 'html', 'htm', 'yml', 'yaml', 'ini', 'conf', 'cfg', 'sql', 'py', 'js', 'java', 'cpp',
                           'c', 'h']

        if file_extension in file_processors:
            content = file_processors[file_extension](file_data, file_name)
            file_type = file_extension
        elif file_extension in text_extensions:
            content = self.extract_text_content(file_data, file_name)
            file_type = 'text'
        else:
            content = {
                "binary_file": True,
                "file_size": len(file_data),
                "file_extension": file_extension,
                "hex_preview": file_data[:64].hex() if len(file_data) > 64 else file_data.hex()
            }
            file_type = 'binary'

        return {
            "file_name": file_name,
            "file_extension": file_extension,
            "file_type": file_type,
            "file_size": len(file_data),
            "content": content
        }

    def send_message_with_confirmation(self, topic, key, value, timeout=30):
        """Отправляет сообщение с подтверждением доставки и повторными попытками"""
        max_retries = 3
        attempt = 0
        while attempt < max_retries:
            try:
                future = self.producer.send(topic, key=key, value=value)
                record_metadata = future.get(timeout=timeout)
                print(
                    f"Сообщение доставлено в {record_metadata.topic} [{record_metadata.partition}] по смещению {record_metadata.offset}")
                return True
            except KafkaTimeoutError as e:
                attempt += 1
                print(f"Попытка {attempt} не удалась из-за таймаута: {e}")
                if attempt < max_retries:
                    print(f"Повторная попытка через {attempt} секунд...")
                    time.sleep(attempt)
                else:
                    print(f"Не удалось отправить сообщение после {max_retries} попыток: {e}")
                    return False
            except KafkaError as e:
                print(f"Ошибка Kafka: {e}")
                return False
            except Exception as e:
                print(f"Непредвиденная ошибка: {e}")
                return False
        return False

    def save_to_postgresql(self, file_name, path_file_s3, upload_timestamp, status='uploaded'):
        """Сохраняет метаданные напрямую в PostgreSQL"""
        try:
            conn = psycopg2.connect(
                host="localhost",
                port="5433",
                database="mydb",
                user="user",
                password="password"
            )

            cursor = conn.cursor()
            insert_query = """
            INSERT INTO file_metadata (file_name, path_file_s3, upload_timestamp, status)
            VALUES (%s, %s, %s, %s)
            """
            cursor.execute(insert_query, (file_name, path_file_s3, upload_timestamp, status))
            conn.commit()

            print(f"Успешно сохранено в PostgreSQL: {file_name} в {upload_timestamp}")
            return True

        except Exception as e:
            print(f"Ошибка сохранения в PostgreSQL: {e}")
            return False
        finally:
            if 'conn' in locals():
                conn.close()

    def process_file(self, file_path, inn):
        """Обрабатывает файл и отправляет его целиком с извлеченным содержимым"""
        file_name = os.path.basename(file_path)
        file_extension = file_name.split('.')[-1] if '.' in file_name else 'bin'

        with open(file_path, 'rb') as f:
            file_data = f.read()

        file_hash = hashlib.md5(file_data).hexdigest()
        file_size = len(file_data)

        # Извлекаем содержимое файла
        extracted_content = self.extract_file_content(file_path)
        parsed_content_json = json.dumps(extracted_content, ensure_ascii=False) if extracted_content else None

        send_id = self.get_next_send_id()
        current_datetime = datetime.now()
        current_date = current_datetime.strftime("%Y-%m-%d")

        print(f"Обработка {file_name} для ИНН {inn}, размер файла: {file_size} байт")
        print(f"Тип файла: {extracted_content['file_type']}")
        print(f"Глобальный ID отправки: {send_id}")

        path_s3 = f"/data/{current_date}/inn_{inn}_send_{send_id}/{file_name}"

        payload = {
            "inn": inn,
            "file_name": file_name,
            "file_extension": file_extension,
            "file_data": file_data.hex(),
            "file_hash": file_hash,
            "file_size": file_size,
            "processing_date": current_date,
            "timestamp": current_datetime.isoformat(),
            "path_s3": path_s3,
            "parsed_content": parsed_content_json,
            "file_type": extracted_content['file_type']
        }

        message = self.create_message_with_schema(payload)
        key = f"{inn}_{current_date}_{send_id}"

        if self.send_message_with_confirmation('file-chunks-topic', key, message):
            print(f"Файл {file_name} успешно отправлен. Сохраняем метаданные в PostgreSQL.")
            if self.send_load_message(inn, send_id, current_datetime, file_name):
                print(f"Файл {file_name} успешно обработан (Глобальный ID отправки: {send_id})")
                return True
            else:
                print(f"Не удалось сохранить метаданные для {file_name}")
                return False
        else:
            print(f"Не удалось отправить файл {file_name}")
            return False

    def process_data_from_api(self, api_data, inn, file_name, file_extension="json"):
        """Обрабатывает данные из API и отправляет в Kafka"""
        try:
            # Преобразуем данные в байты
            if isinstance(api_data, (dict, list)):
                file_data = json.dumps(api_data, ensure_ascii=False, indent=2).encode('utf-8')
            elif isinstance(api_data, str):
                file_data = api_data.encode('utf-8')
            else:
                file_data = str(api_data).encode('utf-8')

            file_hash = hashlib.md5(file_data).hexdigest()
            file_size = len(file_data)

            # Создаем временный файл для обработки содержимого
            with tempfile.NamedTemporaryFile(mode='wb', delete=False, suffix=f'.{file_extension}') as temp_file:
                temp_file.write(file_data)
                temp_file_path = temp_file.name

            try:
                # Извлекаем содержимое файла
                extracted_content = self.extract_file_content(temp_file_path)
                parsed_content_json = json.dumps(extracted_content, ensure_ascii=False) if extracted_content else None
            finally:
                # Удаляем временный файл
                os.unlink(temp_file_path)

            send_id = self.get_next_send_id()
            current_datetime = datetime.now()
            current_date = current_datetime.strftime("%Y-%m-%d")

            print(f"Обработка данных API для ИНН {inn}, размер: {file_size} байт")
            print(f"Тип данных: {file_extension}")
            print(f"Глобальный ID отправки: {send_id}")

            path_s3 = f"/data/{current_date}/inn_{inn}_send_{send_id}/{file_name}"

            payload = {
                "inn": inn,
                "file_name": file_name,
                "file_extension": file_extension,
                "file_data": file_data.hex(),
                "file_hash": file_hash,
                "file_size": file_size,
                "processing_date": current_date,
                "timestamp": current_datetime.isoformat(),
                "path_s3": path_s3,
                "parsed_content": parsed_content_json,
                "file_type": extracted_content['file_type'] if extracted_content else 'api_data'
            }

            message = self.create_message_with_schema(payload)
            key = f"{inn}_{current_date}_{send_id}"

            if self.send_message_with_confirmation('file-chunks-topic', key, message):
                print(f"Данные API {file_name} успешно отправлены. Сохраняем метаданные в PostgreSQL.")
                if self.send_load_message(inn, send_id, current_datetime, file_name):
                    print(f"Данные API {file_name} успешно обработаны (Глобальный ID отправки: {send_id})")
                    return True
                else:
                    print(f"Не удалось сохранить метаданные для {file_name}")
                    return False
            else:
                print(f"Не удалось отправить данные API {file_name}")
                return False

        except Exception as e:
            print(f"Ошибка при обработке данных из API: {e}")
            return False

    def send_load_message(self, inn, send_id, processing_datetime, file_name):
        """Сохраняет метаданные о завершении загрузки напрямую в PostgreSQL"""
        current_date = processing_datetime.strftime("%Y-%m-%d")
        path_file_s3 = f"s3a://moscow-industry-data/topics/file-chunks-topic/path_s3=/data/{current_date}/inn_{inn}_send_{send_id}/{file_name}"

        print(f"Сохранение в PostgreSQL - имя файла: {file_name}, путь: {path_file_s3}, время: {processing_datetime}")

        if self.save_to_postgresql(file_name, path_file_s3, processing_datetime, "uploaded"):
            print(f"Метаданные успешно сохранены в PostgreSQL для ИНН {inn} (Глобальный ID отправки: {send_id})")
            return True
        else:
            print(f"Не удалось сохранить метаданные в PostgreSQL для ИНН {inn} (Глобальный ID отправки: {send_id})")
            return False

    def close(self):
        """Закрывает Kafka продюсер"""
        if self.producer:
            print("Очистка продюсера...")
            self.producer.flush(timeout=30)
            print("Продюсер успешно очищен.")
            self.producer.close()
            print("Продюсер закрыт.")


# Глобальная переменная для продюсера
producer = None


# Функции для работы с PostgreSQL
def check_postgresql_connection():
    """Проверяет подключение к PostgreSQL"""
    try:
        conn = psycopg2.connect(
            host="localhost",
            port="5433",
            database="mydb",
            user="user",
            password="password"
        )
        cursor = conn.cursor()
        cursor.execute("SELECT version();")
        version = cursor.fetchone()
        print(f"Подключение к PostgreSQL успешно. Версия: {version[0]}")
        conn.close()
        return True
    except Exception as e:
        print(f"Ошибка подключения к PostgreSQL: {e}")
        return False


def check_table_exists():
    """Проверяет существование таблицы file_metadata"""
    try:
        conn = psycopg2.connect(
            host="localhost",
            port="5433",
            database="mydb",
            user="user",
            password="password"
        )
        cursor = conn.cursor()
        cursor.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables 
                WHERE table_name = 'file_metadata'
            );
        """)
        exists = cursor.fetchone()[0]
        conn.close()
        return exists
    except Exception as e:
        print(f"Ошибка проверки таблицы: {e}")
        return False


def create_table_if_not_exists():
    """Создает таблицу если она не существует"""
    try:
        conn = psycopg2.connect(
            host="localhost",
            port="5433",
            database="mydb",
            user="user",
            password="password"
        )
        cursor = conn.cursor()
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS file_metadata (
                id SERIAL PRIMARY KEY,
                file_name VARCHAR(255),
                path_file_s3 VARCHAR(500),
                upload_timestamp TIMESTAMP,
                status VARCHAR(50) DEFAULT 'uploaded'
            );
        """)
        conn.commit()
        print("Таблица 'file_metadata' создана или уже существует")
        conn.close()
        return True
    except Exception as e:
        print(f"Ошибка создания таблицы: {e}")
        return False


def get_table_data():
    """Возвращает данные из таблицы file_metadata"""
    try:
        conn = psycopg2.connect(
            host="localhost",
            port="5433",
            database="mydb",
            user="user",
            password="password"
        )
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM file_metadata ORDER BY upload_timestamp DESC;")
        rows = cursor.fetchall()

        data = []
        for row in rows:
            data.append({
                "id": row[0],
                "file_name": row[1],
                "path_file_s3": row[2],
                "upload_timestamp": row[3].isoformat() if row[3] else None,
                "status": row[4]
            })
        conn.close()
        return data
    except Exception as e:
        print(f"Ошибка чтения данных таблицы: {e}")
        return []


def check_superset_availability():
    """Проверяет, доступен ли Superset"""
    try:
        response = requests.get("http://localhost:8088", timeout=5)
        return response.status_code == 200
    except:
        return False


def download_from_api(api_url, timeout=30):
    """Загружает данные из внешнего API"""
    try:
        start_time = time.time()

        # Выполняем GET-запрос
        response = requests.get(api_url, timeout=timeout)
        response.raise_for_status()  # Проверяем статус ответа

        download_time = time.time() - start_time

        # Пытаемся распарсить JSON
        try:
            data = response.json()
            content_type = "json"
        except:
            # Если не JSON, возвращаем как текст
            data = response.text
            content_type = "text"

        return {
            "status": "success",
            "data": data,
            "content_type": content_type,
            "file_size": len(response.content),
            "download_time": download_time,
            "status_code": response.status_code
        }

    except requests.exceptions.Timeout:
        return {
            "status": "error",
            "error": f"Таймаут запроса к API ({timeout} секунд)"
        }
    except requests.exceptions.HTTPError as e:
        return {
            "status": "error",
            "error": f"HTTP ошибка: {e}"
        }
    except requests.exceptions.RequestException as e:
        return {
            "status": "error",
            "error": f"Ошибка запроса: {e}"
        }
    except Exception as e:
        return {
            "status": "error",
            "error": f"Неожиданная ошибка: {e}"
        }


@asynccontextmanager
async def lifespan(app: FastAPI):
    # Запуск: создаем Kafka продюсер
    global producer
    producer = KafkaConnectFileProducer()

    # Проверяем подключение к PostgreSQL при старте
    print("Проверка подключения к PostgreSQL...")
    if not check_postgresql_connection():
        print("Предупреждение: Не удалось подключиться к PostgreSQL")

    print("Проверка существования таблицы...")
    if not check_table_exists():
        print("Создание таблицы...")
        create_table_if_not_exists()

    print("Kafka продюсер запущен")
    yield
    # Завершение: закрываем продюсер
    if producer:
        producer.close()
    print("Kafka продюсер остановлен")


app = FastAPI(
    title="File Processor API",
    description="API для обработки файлов и отправки в Kafka",
    version="1.0.0",
    lifespan=lifespan
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://192.168.3.17:3000",
        "http://192.168.137.1:3000"
    ],
    allow_credentials=True,
    allow_methods=["*"],  # Разрешаем все методы (GET, POST, OPTIONS и т.д.)
    allow_headers=["*"],  # Разрешаем все заголовки
)


def run_sync_processing(file_path: str, inn: str):
    """Вспомогательная функция для запуска синхронного process_file в отдельном потоке."""
    global producer
    if producer is None:
        raise HTTPException(status_code=500, detail="Kafka продюсер недоступен")

    success = producer.process_file(file_path, inn)
    return success


def run_sync_api_processing(api_data: dict, inn: str, file_name: str, file_extension: str):
    """Вспомогательная функция для запуска синхронного process_data_from_api в отдельном потоке."""
    global producer
    if producer is None:
        raise HTTPException(status_code=500, detail="Kafka продюсер недоступен")

    success = producer.process_data_from_api(api_data, inn, file_name, file_extension)
    return success


@app.post("/process-file/", response_model=ProcessResponse)
async def process_file_endpoint(request: FileProcessRequest):
    """
    Обработать файл и отправить в Kafka.

    - **file_path**: Полный путь к файлу на сервере
    - **inn**: ИНН организации (10 или 12 цифр)
    """

    # Проверка входных данных
    if not os.path.exists(request.file_path):
        raise HTTPException(status_code=404, detail="Файл не найден")

    if not request.inn.isdigit() or len(request.inn) not in [10, 12]:
        raise HTTPException(status_code=400, detail="ИНН должен содержать 10 или 12 цифр")

    try:
        # Запускаем синхронную обработку в отдельном потоке
        success = await asyncio.to_thread(run_sync_processing, request.file_path, request.inn)

        if success:
            return ProcessResponse(
                status="success",
                message=f"Файл {os.path.basename(request.file_path)} успешно обработан",
                inn=request.inn,
                file_name=os.path.basename(request.file_path)
            )
        else:
            raise HTTPException(status_code=500, detail="Не удалось обработать файл")

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка обработки: {str(e)}")


@app.post("/download-from-api/", response_model=ApiDownloadResponse)
async def download_from_api_endpoint(request: ApiDownloadRequest):
    """
    Загрузить данные из внешнего API и отправить в Kafka.

    - **api_url**: URL внешнего API
    - **inn**: ИНН организации (10 или 12 цифр)
    - **file_name**: Имя файла (опционально, сгенерируется автоматически если не указано)
    - **file_extension**: Расширение файла (по умолчанию: json)
    """
    # Проверка входных данных
    if not request.inn.isdigit() or len(request.inn) not in [10, 12]:
        raise HTTPException(status_code=400, detail="ИНН должен содержать 10 или 12 цифр")

    if not request.api_url.startswith(('http://', 'https://')):
        raise HTTPException(status_code=400, detail="Некорректный URL API")

    # Генерируем имя файла если не указано
    if not request.file_name:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        request.file_name = f"api_data_{request.inn}_{timestamp}.{request.file_extension}"

    try:
        # Загружаем данные из API
        print(f"Загрузка данных из API: {request.api_url}")
        api_result = download_from_api(request.api_url)

        if api_result["status"] == "error":
            raise HTTPException(status_code=500, detail=f"Ошибка загрузки из API: {api_result['error']}")

        # Запускаем обработку данных в отдельном потоке
        success = await asyncio.to_thread(
            run_sync_api_processing,
            api_result["data"],
            request.inn,
            request.file_name,
            request.file_extension
        )

        if success:
            return ApiDownloadResponse(
                status="success",
                message=f"Данные из API успешно загружены и обработаны",
                inn=request.inn,
                file_name=request.file_name,
                api_url=request.api_url,
                file_size=api_result.get("file_size"),
                download_time=api_result.get("download_time")
            )
        else:
            raise HTTPException(status_code=500, detail="Не удалось обработать данные из API")

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Ошибка обработки: {str(e)}")



@app.get("/open-superset")
async def open_superset(autolaunch: bool = False):
    """
    Открыть Apache Superset для аналитики и визуализации данных

    - **autolaunch**: Автоматически открыть в браузере (работает только локально)
    """
    superset_url = "http://localhost:8088"

    # Проверяем доступность Superset
    is_superset_available = check_superset_availability()

    if autolaunch:
        # Пытаемся открыть в браузере
        try:
            def open_browser():
                time.sleep(1)  # Небольшая задержка для гарантии отправки ответа
                webbrowser.open(superset_url)

            threading.Thread(target=open_browser, daemon=True).start()

            return {
                "status": "success",
                "url": superset_url,
                "message": "Superset открывается в браузере...",
                "superset_available": is_superset_available,
                "credentials": "Логин: admin, пароль: admin"
            }
        except Exception as e:
            return {
                "status": "error",
                "url": superset_url,
                "message": f"Не удалось автоматически открыть браузер: {str(e)}",
                "superset_available": is_superset_available,
                "credentials": "Логин: admin, пароль: admin"
            }
    else:
        # Просто возвращаем ссылку
        return {
            "status": "success",
            "url": superset_url,
            "message": "Перейдите по ссылку для доступа к Superset",
            "superset_available": is_superset_available,
            "credentials": "Логин: admin, пароль: admin"
        }




if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8001)