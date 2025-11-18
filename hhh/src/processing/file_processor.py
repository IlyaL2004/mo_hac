import io
import csv
import pandas as pd
import json
import chardet
import PyPDF2
import docx
from typing import Dict, Any, Optional, Tuple
import re


class FileProcessor:
    """Обработчик файлов различных форматов"""

    """Обработчик файлов различных форматов"""

    def detect_encoding(self, file_bytes: bytes) -> str:
        """Определяет кодировку файла"""
        try:
            result = chardet.detect(file_bytes)
            encoding = result['encoding'] if result['encoding'] else 'utf-8'
            confidence = result['confidence']
            print(f"Определена кодировка: {encoding} (уверенность: {confidence:.2%})")
            return encoding
        except:
            return 'utf-8'

    def extract_inn_from_filename(self, file_name: str) -> str:
        """Извлекает ИНН из имени файла или возвращает дефолтный"""
        if file_name:
            inn_match = re.search(r'\b\d{10,12}\b', file_name)
            if inn_match:
                return inn_match.group(0)
        return "7700000000"



    def extract_csv_content(self, file_data: bytes, file_name: str) -> Optional[Dict[str, Any]]:
        """Извлекает содержимое CSV файла"""
        try:
            encoding = self.detect_encoding(file_data)
            text_content = file_data.decode(encoding)

            csv_reader = csv.reader(io.StringIO(text_content))
            rows = list(csv_reader)

            content = {
                "headers": rows[0] if rows else [],
                "rows": rows[1:] if len(rows) > 1 else [],
                "row_count": len(rows) - 1 if len(rows) > 1 else 0,
                "encoding": encoding
            }

            print(f"CSV файл {file_name}: {content['row_count']} строк")
            return content
        except Exception as e:
            print(f"Ошибка при обработке CSV: {e}")
            return None

    def extract_excel_content(self, file_data: bytes, file_name: str) -> Optional[Dict[str, Any]]:
        """Извлекает содержимое Excel файла"""
        try:
            excel_file = io.BytesIO(file_data)
            excel_data = pd.ExcelFile(excel_file)

            content = {
                "sheets": excel_data.sheet_names,
                "data": {}
            }

            for sheet_name in excel_data.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                content["data"][sheet_name] = {
                    "headers": df.columns.tolist(),
                    "rows": df.fillna('').astype(str).values.tolist(),
                    "row_count": len(df),
                    "columns_count": len(df.columns)
                }

            print(f"Excel файл {file_name}: {len(content['sheets'])} листов")
            return content
        except Exception as e:
            print(f"Ошибка при обработке Excel: {e}")
            return None

    def extract_text_content(self, file_data: bytes, file_name: str) -> Optional[Dict[str, Any]]:
        """Извлекает содержимое текстового файла"""
        try:
            encoding = self.detect_encoding(file_data)
            text_content = file_data.decode(encoding)

            lines = text_content.split('\n')
            content = {
                "line_count": len(lines),
                "char_count": len(text_content),
                "encoding": encoding,
                "preview": text_content[:1000]  # первые 1000 символов
            }

            print(f"Текстовый файл {file_name}: {content['line_count']} строк")
            return content
        except Exception as e:
            print(f"Ошибка при обработке текстового файла: {e}")
            return None

    def extract_json_content(self, file_data: bytes, file_name: str) -> Optional[Dict[str, Any]]:
        """Извлекает содержимое JSON файла"""
        try:
            encoding = self.detect_encoding(file_data)
            text_content = file_data.decode(encoding)

            json_data = json.loads(text_content)
            content = {
                "data": json_data,
                "type": type(json_data).__name__,
                "encoding": encoding
            }

            print(f"JSON файл {file_name}: успешно распарсен")
            return content
        except Exception as e:
            print(f"Ошибка при обработке JSON: {e}")
            return None

    def extract_pdf_content(self, file_data: bytes, file_name: str) -> Optional[Dict[str, Any]]:
        """Извлекает содержимое PDF файла"""
        try:
            pdf_file = io.BytesIO(file_data)
            pdf_reader = PyPDF2.PdfReader(pdf_file)

            text_content = ""
            for page in pdf_reader.pages:
                text_content += page.extract_text() + "\n"

            content = {
                "page_count": len(pdf_reader.pages),
                "text_content": text_content.strip(),
                "char_count": len(text_content)
            }

            print(f"PDF файл {file_name}: {content['page_count']} страниц")
            return content
        except Exception as e:
            print(f"Ошибка при обработке PDF: {e}")
            return None

    def extract_word_content(self, file_data: bytes, file_name: str) -> Optional[Dict[str, Any]]:
        """Извлекает содержимое Word документа"""
        try:
            doc_file = io.BytesIO(file_data)
            doc = docx.Document(doc_file)

            text_content = ""
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content += paragraph.text + "\n"

            content = {
                "paragraph_count": len(doc.paragraphs),
                "text_content": text_content.strip(),
                "char_count": len(text_content)
            }

            print(f"Word файл {file_name}: {content['paragraph_count']} параграфов")
            return content
        except Exception as e:
            print(f"Ошибка при обработке Word документа: {e}")
            return None

    def extract_file_content(self, file_path: str) -> Optional[Dict[str, Any]]:
        """Основная функция для извлечения содержимого файла"""
        import os

        file_name = os.path.basename(file_path)
        file_extension = file_name.split('.')[-1].lower() if '.' in file_name else 'bin'

        with open(file_path, 'rb') as f:
            file_data = f.read()

        file_processors = {
            'csv': self.extract_csv_content,
            'xlsx': self.extract_excel_content,
            'xls': self.extract_excel_content,
            'txt': self.extract_text_content,
            'log': self.extract_text_content,
            'json': self.extract_json_content,
            'pdf': self.extract_pdf_content,
            'docx': self.extract_word_content,
            'doc': self.extract_word_content,
        }

        # Текстовые форматы
        text_extensions = ['xml', 'html', 'htm', 'yml', 'yaml', 'ini', 'conf', 'cfg', 'sql', 'py', 'js', 'java', 'cpp',
                           'c', 'h']

        if file_extension in file_processors:
            content = file_processors[file_extension](file_data, file_name)
            file_type = file_extension
        elif file_extension in text_extensions:
            content = self.extract_text_content(file_data, file_name)
            file_type = 'text'
        else:
            content = {
                "binary_file": True,
                "file_size": len(file_data),
                "file_extension": file_extension,
                "hex_preview": file_data[:64].hex() if len(file_data) > 64 else file_data.hex()
            }
            file_type = 'binary'

        return {
            "file_name": file_name,
            "file_extension": file_extension,
            "file_type": file_type,
            "file_size": len(file_data),
            "content": content
        }